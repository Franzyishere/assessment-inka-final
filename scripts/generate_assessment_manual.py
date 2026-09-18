from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "MANUAL-BOOK-ASSESSMENT-INKA.docx"
LOGO = ROOT / "public" / "images" / "logo" / "logo-inka-sidebar.png"
RED = "C8102E"
DARK = "1F2937"
LIGHT_RED = "FDECEF"
LIGHT_GRAY = "F3F4F6"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=DARK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        shade(table.rows[0].cells[idx], RED)
        set_cell_text(table.rows[0].cells[idx], header, True, WHITE, 9)
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            if row_idx % 2:
                shade(cells[idx], "FAFAFA")
            set_cell_text(cells[idx], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_note(doc, title, text, warning=False):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "FFF4E5" if warning else LIGHT_RED)
    p = cell.paragraphs[0]
    r = p.add_run(f"{title}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string("9A3412" if warning else RED)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph()


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, color in [("Title", 28, RED), ("Heading 1", 18, RED), ("Heading 2", 14, DARK), ("Heading 3", 11, RED)]:
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)

# Header and footer
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_run = header.add_run("MANUAL BOOK  |  ASSESSMENT INKA")
header_run.font.name = "Aptos"
header_run.font.size = Pt(8)
header_run.font.bold = True
header_run.font.color.rgb = RGBColor.from_string(RED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("PT Industri Kereta Api (Persero)  •  Halaman ").font.size = Pt(8)
add_field(footer, "PAGE")

# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(70)
if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(3.3))
doc.add_paragraph().paragraph_format.space_after = Pt(30)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MANUAL BOOK")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(30)
r.font.color.rgb = RGBColor.from_string(RED)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SISTEM ASSESSMENT INKA")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(22)
r.font.color.rgb = RGBColor.from_string(DARK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Panduan Operasional Modul Assessment")
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(107, 114, 128)
doc.add_paragraph().paragraph_format.space_after = Pt(100)
cover_meta = doc.add_table(rows=4, cols=2)
cover_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
cover_meta.style = "Light Shading Accent 1"
for i, (label, value) in enumerate([
    ("Versi Dokumen", "1.0"),
    ("Status", "Final Modul Assessment"),
    ("Tanggal", "26 Agustus 2026"),
    ("Zona Waktu", "WIB (Asia/Jakarta)"),
]):
    set_cell_text(cover_meta.cell(i, 0), label, True, RED)
    set_cell_text(cover_meta.cell(i, 1), value)
doc.add_page_break()

doc.add_heading("Informasi Dokumen", level=1)
add_table(doc, ["Item", "Keterangan"], [
    ["Nama aplikasi", "Sistem Assessment INKA"],
    ["Pemilik proses", "PT Industri Kereta Api (Persero)"],
    ["Ruang lingkup", "Assessment internal"],
    ["Role tercakup", "Super Admin, Admin HCGA, Asesor, Peserta Assessment"],
    ["Klasifikasi", "Dokumen operasional internal"],
], [4.5, 11.5])
add_note(doc, "Catatan", "Kredensial pengguna tidak dicantumkan di dalam manual. Email dan password awal didistribusikan melalui kanal resmi oleh administrator sistem.")

doc.add_heading("Daftar Isi", level=1)
p = doc.add_paragraph()
add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
p.add_run("Klik kanan pada daftar isi lalu pilih Update Field apabila nomor halaman belum muncul.")
doc.add_page_break()

doc.add_heading("1. Pendahuluan", level=1)
doc.add_paragraph("Sistem Assessment INKA adalah aplikasi berbasis web untuk mengelola kegiatan assessment internal secara terpusat. Sistem menghubungkan pengelolaan pengguna, Bank Simulasi, Program Assessment, peserta, tim asesor, jadwal pelaksanaan, pengerjaan simulasi, monitoring, penilaian, dan rekomendasi.")
doc.add_paragraph("Manual ini membahas modul assessment yang telah tersedia pada aplikasi dan digunakan sebagai acuan operasional bagi seluruh role terkait.")

doc.add_heading("1.1 Tujuan Sistem", level=2)
add_bullets(doc, [
    "Memusatkan data pengguna dan peserta assessment.",
    "Menstandarkan materi assessment melalui Bank Simulasi.",
    "Mengelompokkan jadwal, peserta, simulasi, dan asesor dalam Program Assessment.",
    "Mendukung empat rangkaian simulasi assessment internal.",
    "Menyediakan monitoring progres dan hasil secara terkontrol.",
    "Menjaga pembatasan akses berdasarkan tanggung jawab setiap role.",
])

doc.add_heading("1.2 Fitur Utama", level=2)
add_bullets(doc, [
    "Login email dan password dengan dashboard sesuai role.",
    "Manajemen pengguna, role, hak akses, dan audit log.",
    "Pengelolaan peserta assessment, Bank Simulasi, dan Program Assessment.",
    "Penugasan satu atau beberapa asesor untuk seluruh simulasi program.",
    "Jadwal assessment dan status simulasi peserta.",
    "Preview materi PDF, jawaban satu kolom atau per materi, dan upload presentasi PDF.",
    "Mode fullscreen wajib saat pengerjaan simulasi bertimer.",
    "Monitoring admin dan asesor, penilaian draft/final, serta rekomendasi peserta.",
])

doc.add_heading("2. Peran dan Hak Akses", level=1)
add_table(doc, ["Role", "Tanggung Jawab"], [
    ["Super Admin", "Mengelola akun seluruh pengguna, informasi role/hak akses, dan audit log."],
    ["Admin HCGA", "Mengelola peserta, Bank Simulasi, Program Assessment, penugasan asesor, dan monitoring."],
    ["Asesor", "Memantau peserta yang ditugaskan serta memberi penilaian dan rekomendasi."],
    ["Peserta Assessment", "Melihat jadwal, mengerjakan simulasi, mengunggah presentasi, dan melihat hasil final."],
], [4.2, 11.8])
add_note(doc, "Keamanan akses", "Menu yang tidak sesuai role tidak ditampilkan. Sistem juga melindungi route sehingga alamat halaman role lain tidak dapat dibuka secara langsung.")

doc.add_heading("3. Akses Sistem", level=1)
doc.add_heading("3.1 Login", level=2)
add_steps(doc, [
    "Buka alamat aplikasi melalui browser yang didukung.",
    "Masukkan email akun yang telah didaftarkan.",
    "Masukkan password.",
    "Opsional: aktifkan Ingat Saya pada perangkat pribadi yang aman.",
    "Tekan Masuk. Sistem mengarahkan pengguna ke dashboard sesuai role.",
])
doc.add_heading("3.2 Logout", level=2)
add_steps(doc, ["Buka menu profil di kanan atas.", "Pilih Keluar.", "Pastikan halaman kembali ke layar login, terutama saat menggunakan perangkat bersama."])
doc.add_heading("3.3 Rekomendasi Perangkat", level=2)
add_bullets(doc, ["Gunakan Google Chrome atau Microsoft Edge versi terbaru.", "Gunakan koneksi internet stabil.", "Peserta disarankan memakai laptop/desktop dan menutup aplikasi yang tidak diperlukan.", "Pastikan browser mengizinkan mode fullscreen dan pembukaan PDF."])

doc.add_page_break()
doc.add_heading("4. Panduan Super Admin", level=1)
doc.add_heading("4.1 Manajemen Pengguna", level=2)
add_steps(doc, [
    "Pilih Manajemen Pengguna.",
    "Tekan Tambah Pengguna.",
    "Isi nama, email, role, dan password awal.",
    "Simpan data dan distribusikan kredensial secara terpisah.",
    "Gunakan Edit untuk memperbarui identitas, role, status, atau password pengguna.",
])
doc.add_heading("4.2 Role dan Hak Akses", level=2)
doc.add_paragraph("Menu Role & Hak Akses menampilkan kewenangan masing-masing pengguna. Perubahan role harus dilakukan dengan hati-hati karena menentukan dashboard, menu, dan route yang dapat diakses.")
doc.add_heading("4.3 Audit Log", level=2)
doc.add_paragraph("Audit Log digunakan untuk meninjau aktivitas penting seperti login, logout, perubahan pengguna, pengaturan program, pilihan Simulasi 3 peserta Madya, dan aktivitas administratif lain yang telah dicatat sistem.")

doc.add_heading("5. Panduan Admin HCGA", level=1)
doc.add_heading("5.1 Mengelola Akun Peserta", level=2)
add_steps(doc, [
    "Pilih Peserta Assessment.",
    "Tekan Tambah Peserta Assessment.",
    "Isi nama, email, password awal, dan status akun.",
    "Simpan akun.",
    "Gunakan pencarian untuk menemukan peserta, lalu pilih Edit atau Hapus sesuai kebutuhan.",
])
add_note(doc, "Validasi penghapusan", "Peserta yang telah memiliki keterkaitan dengan Program Assessment tidak dapat dihapus agar riwayat tetap terjaga.", True)

doc.add_heading("5.2 Mengelola Bank Simulasi", level=2)
doc.add_paragraph("Jenis simulasi telah disediakan sistem. Admin tidak perlu membuat jenis simulasi baru; admin cukup memperbarui materi PDF dan durasi sesuai kebutuhan pelaksanaan.")
add_table(doc, ["Urutan", "Simulasi", "Mekanisme"], [
    ["1", "Problem Analysis (PA)", "Satu PDF untuk seluruh peserta dan satu kolom jawaban penuh."],
    ["2", "Leaderless Group Discussion (LGD)", "Menampilkan materi dan jawaban PA secara read-only, memakai timer, lalu peserta menyelesaikan sesi."],
    ["3", "Critical Incident / In-Tray", "Beberapa PDF; setiap materi memiliki kolom jawaban; satu timer untuk seluruh paket."],
    ["4", "Presentasi", "Peserta mengunggah PDF; asesor dapat preview dan fullscreen."],
], [1.5, 5, 9.5])
add_steps(doc, [
    "Buka Bank Simulasi.",
    "Pilih simulasi atau paket materi yang akan diperbarui.",
    "Unggah PDF yang benar dan tentukan durasi bila tersedia.",
    "Simpan perubahan dan periksa kembali nama file/materi yang ditampilkan.",
])

doc.add_heading("5.3 Paket Simulasi 3", level=2)
add_table(doc, ["Paket", "Kategori Peserta"], [
    ["Critical Incident 1", "Golongan I ke II, Level Spesialis Pratama, Level SPV"],
    ["Critical Incident 2", "Golongan II ke III, Level Spesialis Muda"],
    ["Critical Incident 3", "Golongan III ke IV; salah satu pilihan Spesialis Madya"],
    ["In-Tray 1", "Level M"],
    ["In-Tray 2", "Level SM"],
    ["In-Tray 3", "Salah satu pilihan Spesialis Madya"],
], [5, 11])

doc.add_heading("5.4 Membuat Program Assessment", level=2)
add_steps(doc, [
    "Pilih Program Assessment dan tekan Buat Program.",
    "Isi nama program, deskripsi, waktu mulai, waktu selesai, dan status.",
    "Simpan program.",
    "Buka Atur Program untuk memilih peserta, kategori, dan tim asesor.",
])
add_table(doc, ["Status", "Penggunaan"], [
    ["Draft", "Program masih dipersiapkan dan aman untuk disunting."],
    ["Aktif", "Program tersedia untuk pelaksanaan sesuai jadwal."],
    ["Arsip", "Program disimpan sebagai riwayat dan tidak digunakan untuk sesi baru."],
], [4, 12])
add_note(doc, "Kode program", "Kode program tidak ditampilkan atau diminta pada antarmuka pengguna. Sistem dapat menggunakan identitas internal untuk kebutuhan teknis tanpa membebani pengguna.")

doc.add_heading("5.5 Menyusun Program dan Tim Asesor", level=2)
add_steps(doc, [
    "Buka Atur Program pada program yang dipilih.",
    "Cari dan pilih peserta yang mengikuti kegiatan.",
    "Tentukan kategori assessment untuk setiap peserta.",
    "Pilih minimal satu asesor atau gunakan pilihan semua asesor bila diperlukan.",
    "Simpan pengaturan. Sistem memasang simulasi dan menerapkan tim asesor ke seluruh simulasi program.",
])
add_note(doc, "Validasi", "Program tidak dapat disimpan tanpa asesor. Satu program dapat berisi peserta dari seluruh kategori pada hari yang sama.", True)

doc.add_heading("5.6 Penugasan Asesor", level=2)
doc.add_paragraph("Penugasan asesor berlaku untuk seluruh simulasi dalam satu program, bukan dipilih satu per satu per simulasi. Gunakan menu Penugasan Asesor untuk memperbarui tim program. Sistem memastikan akun yang dipilih benar-benar memiliki role Asesor.")

doc.add_heading("5.7 Monitoring", level=2)
add_bullets(doc, ["Jumlah peserta dan target sesi.", "Sesi yang telah dimulai.", "Submission yang telah masuk.", "Penilaian yang telah difinalisasi.", "Aktivitas fullscreen yang tercatat.", "Progres per peserta dan simulasi."])

doc.add_page_break()
doc.add_heading("6. Panduan Asesor", level=1)
doc.add_heading("6.1 Simulasi Ditugaskan", level=2)
add_steps(doc, [
    "Pilih Simulasi Ditugaskan.",
    "Periksa jumlah submission dan progress bar.",
    "Pilih Lihat Peserta untuk membuka daftar peserta.",
    "Buka preview jawaban atau PDF jika submission telah tersedia.",
])
doc.add_paragraph("Simulasi aktif disembunyikan dari daftar setelah seluruh peserta relevan menyelesaikannya. Riwayat yang siap dinilai tetap tersedia pada Penilaian & Rekomendasi.")

doc.add_heading("6.2 Peserta Assessment", level=2)
doc.add_paragraph("Menu ini menampilkan peserta dari program yang menjadi tanggung jawab asesor. Gunakan kolom pencarian nama/email untuk data berjumlah besar. Progress menunjukkan simulasi yang dimulai dan dikumpulkan.")

doc.add_heading("6.3 Monitoring Simulasi", level=2)
add_bullets(doc, ["Persentase pengumpulan.", "Jumlah peserta, sesi mulai, submission masuk, dan penilaian.", "Aktivitas fullscreen yang tercatat.", "Akses langsung ke detail peserta."])

doc.add_heading("6.4 Penilaian dan Rekomendasi", level=2)
add_steps(doc, [
    "Pilih Penilaian & Rekomendasi.",
    "Buka submission dengan tombol Beri Penilaian.",
    "Periksa materi PDF, jawaban peserta, atau presentasi PDF.",
    "Isi catatan assessment.",
    "Pilih rekomendasi.",
    "Pilih Simpan Draft apabila belum final, atau Finalisasi untuk mengunci hasil.",
])
add_table(doc, ["Rekomendasi", "Makna"], [
    ["Direkomendasikan", "Peserta memenuhi rekomendasi berdasarkan hasil assessment."],
    ["Direkomendasikan dengan Pengembangan", "Peserta direkomendasikan dengan catatan area pengembangan."],
    ["Belum Direkomendasikan", "Peserta belum direkomendasikan berdasarkan hasil yang tersedia."],
], [6, 10])
add_note(doc, "Finalisasi", "Penilaian final tidak dapat diubah melalui alur normal. Periksa rekomendasi dan catatan sebelum finalisasi.", True)

doc.add_heading("7. Panduan Peserta Assessment", level=1)
doc.add_heading("7.1 Simulasi Saya", level=2)
add_steps(doc, [
    "Pilih Simulasi Saya.",
    "Periksa program, kategori, jadwal, durasi, dan status simulasi.",
    "Tekan Lihat Detail untuk membaca instruksi.",
    "Tekan Mulai Simulasi hanya ketika siap mengerjakan.",
])
doc.add_paragraph("Simulasi yang sudah selesai atau melewati waktu pelaksanaan otomatis disembunyikan agar daftar tetap bersih.")

doc.add_heading("7.2 Jadwal Assessment", level=2)
doc.add_paragraph("Menu Jadwal Assessment menampilkan program dan seluruh simulasi yang relevan dengan kategori peserta, termasuk waktu mulai, waktu berakhir, durasi, dan status.")
add_table(doc, ["Status Jadwal", "Keterangan"], [
    ["Akan Datang", "Waktu buka simulasi belum tiba."],
    ["Tersedia", "Simulasi sudah dapat dibuka."],
    ["Sedang Dikerjakan", "Sesi telah dimulai dan belum dikumpulkan."],
    ["Selesai", "Submission telah dikumpulkan."],
    ["Jadwal Berakhir", "Batas waktu simulasi atau program telah lewat."],
], [5, 11])

doc.add_heading("7.3 Ketentuan Sebelum Memulai", level=2)
add_bullets(doc, ["Pastikan koneksi stabil dan perangkat memiliki daya cukup.", "Siapkan waktu sesuai durasi simulasi.", "Gunakan mode fullscreen selama pengerjaan.", "Jangan menekan tombol kumpulkan sebelum jawaban diperiksa.", "Submission yang telah dikumpulkan tidak dapat diubah."])

doc.add_heading("7.4 Simulasi 1 — Problem Analysis", level=2)
add_steps(doc, ["Mulai Simulasi 1.", "Aktifkan fullscreen.", "Baca PDF pada preview besar.", "Isi satu kolom jawaban penuh.", "Pilih Simpan & Kumpulkan. Sistem menolak jawaban wajib yang kosong."])

doc.add_heading("7.5 Simulasi 2 — Leaderless Group Discussion", level=2)
add_steps(doc, ["Simulasi 1 harus selesai terlebih dahulu.", "Mulai Simulasi 2 dan aktifkan fullscreen.", "Review PDF dan jawaban Simulasi 1 dalam kondisi read-only sebagai bahan LGD.", "Ikuti kegiatan LGD dan perhatikan timer.", "Centang konfirmasi lalu pilih Simulasi Sudah Selesai."])
add_note(doc, "Proses offline", "Laporan LGD dilaksanakan secara offline dan tidak diunggah melalui aplikasi.")

doc.add_heading("7.6 Simulasi 3 — Critical Incident / In-Tray", level=2)
add_steps(doc, ["Sistem menampilkan paket sesuai kategori peserta.", "Khusus Spesialis Madya, pilih CI 3 atau In-Tray 3 dan konfirmasi pilihan final.", "Mulai simulasi dan aktifkan fullscreen.", "Baca setiap materi PDF dan isi jawaban tepat di bawah materi.", "Perhatikan satu timer yang berlaku untuk seluruh paket.", "Periksa seluruh jawaban lalu kumpulkan."])
add_note(doc, "Pilihan Madya", "Pilihan CI 3 atau In-Tray 3 dikunci setelah dikonfirmasi dan tidak dapat diganti.", True)

doc.add_heading("7.7 Simulasi 4 — Presentasi", level=2)
add_steps(doc, ["Buka Simulasi Presentasi.", "Pilih file PDF final.", "Pastikan ukuran file maksimal 25 MB.", "Tekan Kumpulkan Presentasi.", "Pastikan notifikasi berhasil tampil."])
doc.add_paragraph("Asesor dapat membuka presentasi melalui preview besar dan mode fullscreen.")

doc.add_heading("7.8 Hasil & Rekomendasi", level=2)
doc.add_paragraph("Hasil hanya tampil setelah asesor memfinalisasi penilaian. Peserta dapat melihat rekomendasi, nama asesor, tanggal penilaian, dan catatan assessment. Draft asesor tidak dipublikasikan.")

doc.add_heading("8. Timer dan Mode Fullscreen", level=1)
doc.add_heading("8.1 Timer", level=2)
add_bullets(doc, ["Timer dimulai setelah peserta menekan tombol mulai.", "Simulasi 3 menggunakan satu timer untuk seluruh materi.", "Indikator berubah menjadi kuning saat waktu menipis dan merah pada kondisi kritis.", "Peserta harus menyimpan dan mengumpulkan sebelum waktu berakhir."])
doc.add_heading("8.2 Fullscreen Wajib", level=2)
doc.add_paragraph("Pada halaman pengerjaan, peserta wajib masuk fullscreen. Jika keluar, halaman dikunci sampai fullscreen diaktifkan kembali. Aktivitas tersebut dicatat sebagai konteks monitoring, bukan keputusan otomatis terhadap hasil peserta.")

doc.add_heading("9. Matriks Simulasi", level=1)
matrix = [
    ["Kenaikan Golongan I ke II", "Ya", "Ya", "CI 1", "—", "Ya"],
    ["Kenaikan Golongan II ke III", "Ya", "Ya", "CI 2", "—", "Ya"],
    ["Kenaikan Golongan III ke IV", "Ya", "Ya", "CI 3", "—", "Ya"],
    ["Level SPV", "Ya", "Ya", "CI 1", "—", "Ya"],
    ["Level Spesialis Pratama", "Ya", "Ya", "CI 1", "—", "Ya"],
    ["Level M", "Ya", "Ya", "—", "In-Tray 1", "Ya"],
    ["Level Spesialis Muda", "Ya", "Ya", "CI 2", "—", "Ya"],
    ["Level SM", "Ya", "Ya", "—", "In-Tray 2", "Ya"],
    ["Level Spesialis Madya", "Ya", "Ya", "Pilih CI 3", "atau In-Tray 3", "Ya"],
]
add_table(doc, ["Kategori", "PA", "LGD", "CI", "In-Tray", "Presentasi"], matrix, [5.2, 1.2, 1.2, 2.8, 3, 2.4])
add_note(doc, "Catatan", "Wawancara setelah Critical Incident atau In-Tray dilakukan secara offline dan tidak menjadi input aplikasi.")

doc.add_heading("10. Skenario Operasional End-to-End", level=1)
add_steps(doc, [
    "Super Admin memastikan akun Admin HCGA dan asesor aktif.",
    "Admin HCGA membuat akun peserta assessment.",
    "Admin memperbarui materi PA, paket CI/In-Tray, dan durasi pada Bank Simulasi.",
    "Admin membuat Program Assessment berstatus Draft dan menentukan periode.",
    "Admin membuka Atur Program, memilih seluruh peserta, menentukan kategori, dan memilih tim asesor.",
    "Admin memeriksa jadwal dan mengaktifkan program.",
    "Peserta membuka Jadwal Assessment dan Simulasi Saya.",
    "Peserta menyelesaikan PA, LGD, paket Simulasi 3, dan upload presentasi.",
    "Admin dan asesor memonitor progres.",
    "Asesor membuka submission, menyimpan draft bila perlu, lalu memfinalisasi rekomendasi.",
    "Peserta membuka Hasil & Rekomendasi untuk melihat hasil final.",
])

doc.add_heading("11. Checklist Sebelum Pelaksanaan", level=1)
add_bullets(doc, [
    "Akun Admin HCGA, asesor, dan peserta aktif.",
    "Materi PA dan semua paket CI/In-Tray sudah benar.",
    "File PDF dapat dibuka dan tidak rusak.",
    "Durasi setiap simulasi sudah dikonfirmasi.",
    "Program, periode, dan status sudah benar.",
    "Kategori seluruh peserta sudah dipilih.",
    "Minimal satu asesor telah ditugaskan.",
    "Jadwal peserta tampil sesuai rencana.",
    "Koneksi dan perangkat peserta siap.",
])

doc.add_heading("12. Validasi dan Aturan Penting", level=1)
add_table(doc, ["Area", "Aturan Sistem"], [
    ["Peserta", "Akun yang sudah memiliki riwayat program tidak dapat dihapus sembarangan."],
    ["Program", "Penghapusan dibatasi pada program Draft tanpa riwayat pengerjaan."],
    ["Tim asesor", "Tidak boleh kosong dan hanya akun role Asesor yang dapat dipilih."],
    ["Simulasi 1/3", "Jawaban wajib harus terisi sebelum pengumpulan."],
    ["Simulasi 2", "PA harus selesai dan peserta wajib mengonfirmasi penyelesaian."],
    ["Simulasi 3 Madya", "Pilihan paket bersifat final."],
    ["Presentasi", "Hanya PDF maksimal 25 MB."],
    ["Penilaian", "Rekomendasi wajib saat finalisasi dan hasil final tidak dapat diedit melalui alur normal."],
    ["Hasil peserta", "Hanya penilaian final milik peserta yang sedang login yang ditampilkan."],
], [4.2, 11.8])

doc.add_heading("13. Troubleshooting", level=1)
add_table(doc, ["Kendala", "Tindakan"], [
    ["Tidak dapat login", "Periksa email/password, status akun, Caps Lock, dan hubungi administrator bila tetap gagal."],
    ["Simulasi tidak muncul", "Periksa penugasan program, status program, kategori peserta, serta waktu buka/tutup."],
    ["Tidak dapat memulai LGD", "Pastikan Simulasi 1 telah dikumpulkan."],
    ["PDF tidak tampil", "Muat ulang halaman, izinkan PDF di browser, atau buka melalui tab baru."],
    ["Fullscreen keluar", "Tekan tombol masuk fullscreen kembali untuk membuka halaman pengerjaan."],
    ["Upload presentasi gagal", "Pastikan format PDF, ukuran maksimal 25 MB, koneksi stabil, dan file tidak rusak."],
    ["Hasil belum tampil", "Pastikan asesor sudah melakukan finalisasi, bukan hanya menyimpan draft."],
    ["Jadwal tidak sesuai", "Admin perlu memeriksa periode program dan waktu setiap simulasi."],
], [5.2, 10.8])

doc.add_heading("14. Penutup", level=1)
doc.add_paragraph("Manual Book Sistem Assessment INKA versi 1.0 menjadi panduan operasional modul assessment yang tersedia pada aplikasi per 26 Agustus 2026. Perubahan proses bisnis, jenis simulasi, matriks kategori, atau kebijakan penilaian perlu diikuti pembaruan dokumen agar tetap sesuai dengan sistem.")

# Document properties and final formatting
doc.core_properties.title = "Manual Book Sistem Assessment INKA"
doc.core_properties.subject = "Panduan operasional modul assessment"
doc.core_properties.author = "PT Industri Kereta Api (Persero)"
doc.core_properties.keywords = "INKA, assessment, manual book, admin, asesor, peserta"

for sec in doc.sections:
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
