from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\Oddy\Documents\INKA\Document\PROPOSAL & BLUEPRINT COUNSELING CORNER.docx")
OUTPUT = ROOT / "docs" / "PROPOSAL-BLUEPRINT-INKA-ASSESSMENT-MANAGEMENT-SYSTEM-REVISI.docx"
LOGO = ROOT / "public" / "images" / "logo" / "logo-inka-sidebar.png"

FONT = "Times New Roman"
RED = "CC0000"
DARK = "222222"
GRAY = "595959"
LIGHT_GRAY = "F2F2F2"
LIGHT_RED = "FCE8E8"
WHITE = "FFFFFF"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_table_borders(table):
    table_properties = table._tbl.tblPr
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "BFBFBF")


def shade(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell(cell, text, *, bold=False, color=DARK, size=10, center=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], RED)
        set_cell(table.rows[0].cells[index], header, bold=True, color=WHITE, center=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2:
                shade(cells[index], "FAFAFA")
            set_cell(cells[index], value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph()
    return table


def add_field(paragraph, instruction: str, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_heading(doc, text: str, level=1, *, centered=False):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK if level > 1 else "000000")
    return paragraph


def add_paragraph(doc, text: str, *, bold_prefix: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.3)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        first, rest = text[: len(bold_prefix)], text[len(bold_prefix) :]
        run = paragraph.add_run(first)
        run.bold = True
        run.font.name = FONT
        run.font.size = Pt(11)
        run = paragraph.add_run(rest)
    else:
        run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(DARK)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = Inches(0.3)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(f"•  {item}")
        run.font.name = FONT
        run.font.size = Pt(11)


def add_steps(doc, items):
    for number, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = Inches(0.3)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(f"{number}.  {item}")
        run.font.name = FONT
        run.font.size = Pt(11)


def add_callout(doc, title: str, text: str, *, warning=False):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "FFF2CC" if warning else LIGHT_RED)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(f"{title}: ")
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("9C6500" if warning else RED)
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(DARK)
    doc.add_paragraph()


def clear_body_keep_section(document):
    body = document._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def copy_page_setup(source_document, target_document):
    target = target_document.sections[0]
    # Template menggunakan kertas Letter dan margin 1 inci. Nilai dibuat
    # eksplisit agar tidak terjadi konversi ganda EMU -> twip saat disalin.
    target.page_width = Inches(8.5)
    target.page_height = Inches(11)
    target.top_margin = Inches(1)
    target.bottom_margin = Inches(1)
    target.left_margin = Inches(1)
    target.right_margin = Inches(1)
    target.header_distance = Inches(0.5)
    target.footer_distance = Inches(0.5)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(12)
        style.font.bold = True


if not TEMPLATE.exists():
    raise FileNotFoundError(f"Template tidak ditemukan: {TEMPLATE}")

template_document = Document(TEMPLATE)
# Bangun package OOXML baru agar relationship/gambar floating dari template lama
# tidak ikut terbawa. Format halaman dan gaya visual tetap mengacu pada template.
doc = Document()
clear_body_keep_section(doc)
copy_page_setup(template_document, doc)
configure_styles(doc)

# Header/footer mengikuti karakter template: header kosong dan nomor halaman di tengah.
section = doc.sections[0]
header = section.header.paragraphs[0]
header.text = ""
footer = section.footer.paragraphs[0]
footer.text = ""
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Halaman ")
footer_run.font.name = FONT
footer_run.font.size = Pt(9)
add_field(footer, "PAGE", "1")

# Cover
cover_space = doc.add_paragraph()
cover_space.paragraph_format.space_after = Pt(40)
if LOGO.exists():
    logo_paragraph = doc.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.add_run().add_picture(str(LOGO), width=Inches(3.25))
    logo_paragraph.paragraph_format.space_after = Pt(28)

for text, size, color in [
    ("PROPOSAL & BLUEPRINT SISTEM", 26, "000000"),
    ("INKA Assessment Management System", 26, RED),
    ("Platform Pengelolaan Assessment Internal Berbasis Digital", 13, GRAY),
    ("PT INKA (Persero)", 16, "000000"),
    ("Divisi Human Capital", 13, "000000"),
]:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5 if size != 13 else 12)
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = size in (26, 16)
    run.font.color.rgb = RGBColor.from_string(color)

doc.add_paragraph().paragraph_format.space_after = Pt(38)
cover_table = doc.add_table(rows=3, cols=2)
cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for row_index, (label, value) in enumerate(
    [("Versi Dokumen", "1.0"), ("Tahun", "2026"), ("Status", "Draft untuk Review")]
):
    set_cell(cover_table.cell(row_index, 0), label, bold=True)
    set_cell(cover_table.cell(row_index, 1), value)
doc.add_page_break()

# Daftar isi
add_heading(doc, "DAFTAR ISI", level=1, centered=True)
toc = doc.add_paragraph()
add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "Klik kanan lalu pilih Update Field untuk memperbarui daftar isi.")
doc.add_page_break()

# Executive summary
add_heading(doc, "EXECUTIVE SUMMARY", level=1, centered=True)
add_paragraph(
    doc,
    "INKA Assessment Management System merupakan aplikasi berbasis web untuk mengelola pelaksanaan assessment internal PT INKA (Persero) secara terpusat. Sistem menghubungkan pengelolaan program, katalog simulasi, peserta, tim asesor, jadwal, pengerjaan, monitoring, penilaian, dan penyampaian hasil dalam satu alur yang memiliki pembatasan akses berdasarkan peran pengguna.",
)
add_paragraph(
    doc,
    "Blueprint ini disusun berdasarkan source code aplikasi yang tersedia pada 7 September 2026. Fitur yang dinyatakan aktif adalah fitur yang ditemukan pada route, controller, model, migration, antarmuka, dan pengujian aplikasi. Rencana seperti sinkronisasi HRIS, Single Sign-On INKA, serta undangan melalui email dipisahkan secara eksplisit agar tidak dipersepsikan sebagai kemampuan produksi yang sudah tersedia.",
)
add_paragraph(
    doc,
    "Sistem mendukung empat rangkaian simulasi: Problem Analysis, Leaderless Group Discussion, Critical Incident/In-Tray, dan Presentasi. Materi assessment disimpan sebagai PDF privat, jawaban tertulis menggunakan rich text editor, timer dipertahankan pada sisi server, dan aktivitas fullscreen tertentu dicatat sebagai konteks monitoring. Asesor dapat memberikan rekomendasi per submission, sedangkan peserta hanya dapat melihat hasil yang telah difinalisasi.",
)
add_callout(
    doc,
    "Nilai Strategis",
    "Platform membentuk satu sumber data operasional assessment, mengurangi pengaturan berulang, memberi visibilitas progres kepada Admin HCGA dan asesor, serta menjaga pemisahan akses antara pengelola, penilai, dan peserta.",
)

# I
add_heading(doc, "I. Latar Belakang Permasalahan", level=1)
add_paragraph(
    doc,
    "Pelaksanaan assessment internal melibatkan banyak komponen yang saling bergantung: jadwal kegiatan, kategori peserta, materi simulasi, penugasan lebih dari satu asesor, durasi pengerjaan, submission peserta, dan hasil penilaian. Tanpa sistem terpusat, konsistensi konfigurasi dan keterlacakan progres berisiko bergantung pada koordinasi manual antarpihak.",
)
add_paragraph(
    doc,
    "INKA Assessment Management System dirancang untuk menempatkan Program Assessment sebagai wadah pelaksanaan. Satu program dapat memuat peserta dari beberapa kategori pada hari yang sama, menggunakan katalog simulasi yang konsisten, serta menerapkan satu tim asesor ke seluruh rangkaian simulasi program.",
)
add_paragraph(doc, "Kebutuhan operasional yang dijawab oleh implementasi aplikasi meliputi:")
add_bullets(
    doc,
    [
        "Pemusatan data program, jadwal, peserta, kategori assessment, dan penugasan asesor.",
        "Pengelolaan materi PDF yang konsisten melalui Bank Simulasi tetap.",
        "Pengerjaan bertimer yang tidak kembali ke awal ketika halaman dimuat ulang atau perangkat terputus.",
        "Penyimpanan jawaban panjang dengan fasilitas pemformatan dan tabel.",
        "Monitoring progres serta event fullscreen sebagai konteks pengawasan.",
        "Penilaian draft/final dan publikasi hasil yang terkontrol kepada peserta.",
    ],
)
add_callout(
    doc,
    "Batasan ruang lingkup",
    "Dokumen ini hanya membahas modul Assessment. Sistem Recruitment tidak termasuk dalam proposal ini dan telah dipisahkan dari aplikasi aktif.",
)

# II
add_heading(doc, "II. Tujuan & Manfaat Sistem", level=1)
add_heading(doc, "Tujuan Pengembangan", level=2)
add_bullets(
    doc,
    [
        "Menyediakan platform operasional assessment internal yang terstruktur dan berbasis role.",
        "Menstandarkan empat rangkaian simulasi dan paket materi Simulasi 3.",
        "Menjaga kesinambungan session, timer, submission, dan hasil penilaian.",
        "Memberikan monitoring program secara ringkas kepada Admin HCGA dan asesor.",
        "Menyediakan landasan teknis untuk integrasi identitas HRIS dan SSO pada tahap berikutnya.",
    ],
)
add_heading(doc, "Manfaat Sistem", level=2)
add_heading(doc, "Bagi Admin HCGA", level=3)
add_table(
    doc,
    ["Aspek", "Manfaat yang Diperoleh"],
    [
        ["Pengelolaan terpusat", "Program, peserta, simulasi, dan tim asesor dikelola dalam satu aplikasi."],
        ["Efisiensi konfigurasi", "Tim asesor diterapkan untuk seluruh simulasi program dan materi yang sama tidak perlu diinput berulang."],
        ["Monitoring", "Progres mulai, submission, penilaian, dan event fullscreen dapat dilihat per program."],
        ["Kontrol data", "Validasi membatasi penghapusan atau perubahan data yang telah memiliki riwayat tertentu."],
    ],
)
add_heading(doc, "Bagi Asesor", level=3)
add_table(
    doc,
    ["Aspek", "Manfaat yang Diperoleh"],
    [
        ["Workspace program", "Simulasi, peserta, monitoring, dan penilaian dikelompokkan berdasarkan program."],
        ["Akses submission", "Materi, jawaban, dan presentasi peserta dapat diperiksa sesuai penugasan."],
        ["Penilaian terkendali", "Review dapat disimpan sebagai draft sebelum difinalisasi."],
        ["Paket Madya", "Asesor menentukan CI 3 atau In-Tray 3 sebelum peserta memulai Simulasi 3."],
    ],
)
add_heading(doc, "Bagi Peserta", level=3)
add_table(
    doc,
    ["Aspek", "Manfaat yang Diperoleh"],
    [
        ["Kejelasan pelaksanaan", "Jadwal, status, durasi, dan simulasi yang relevan ditampilkan berdasarkan penugasan."],
        ["Kesinambungan session", "Timer tetap melanjutkan waktu server ketika peserta kembali ke session aktif."],
        ["Kemudahan menjawab", "Rich text editor mendukung format teks dan tabel untuk jawaban yang kompleks."],
        ["Kerahasiaan hasil", "Peserta hanya melihat hasil miliknya yang telah difinalisasi asesor."],
    ],
)

# III
add_heading(doc, "III. Spesifikasi Teknis", level=1)
add_heading(doc, "Technology Stack", level=2)
add_table(
    doc,
    ["Komponen", "Teknologi", "Versi/Status"],
    [
        ["Backend Framework", "Laravel", "12.x"],
        ["Bahasa Backend", "PHP", "8.2 atau lebih tinggi"],
        ["View Rendering", "Laravel Blade", "Server-side rendering"],
        ["CSS Framework", "Tailwind CSS", "4.1.x"],
        ["Interaktivitas", "Alpine.js", "3.14.x"],
        ["Rich Text Editor", "Tiptap", "3.31.x"],
        ["Visualisasi", "ApexCharts", "5.3.x"],
        ["Build Tool", "Vite", "7.x"],
        ["Automated Test", "Pest + PHPUnit", "Pest 4"],
        ["Database", "Relational database", "Dikonfigurasi melalui environment"],
    ],
)
add_heading(doc, "Arsitektur Aplikasi", level=2)
add_table(
    doc,
    ["Lapisan", "Komponen", "Tanggung Jawab"],
    [
        ["Presentation", "Blade, Tailwind, Alpine.js", "Menampilkan halaman role, form, tabel, modal, toast, dan editor."],
        ["Routing & Security", "Laravel Route, auth, role middleware", "Memisahkan endpoint Super Admin, Admin, Asesor, dan Peserta."],
        ["Application", "Controller dan Form Request", "Menjalankan use case, validasi, transaksi, dan pemeriksaan kepemilikan data."],
        ["Domain/Data", "Eloquent Model", "Mengelola program, katalog, assignment, session, submission, review, dan audit."],
        ["Persistence", "Database dan filesystem", "Menyimpan data relasional serta file PDF privat."],
    ],
)
add_paragraph(
    doc,
    "Aplikasi menggunakan arsitektur monolith Laravel. Halaman utama dirender di server, sedangkan AJAX dipakai secara selektif untuk pencarian realtime, penyimpanan jawaban antarmateri, event fullscreen, dan notifikasi header. Pendekatan ini menjaga implementasi tetap sederhana sekaligus mengurangi pemuatan JavaScript yang tidak diperlukan pada setiap halaman.",
)
add_heading(doc, "Blueprint Data", level=2)
add_table(
    doc,
    ["Domain", "Entitas Utama", "Relasi Operasional"],
    [
        ["Identitas", "users", "Akun dan role pengguna; menampung field kesiapan HRIS."],
        ["Katalog", "simulation_types, simulation_scenarios, simulation_material_pages", "Jenis simulasi, paket, durasi, dan materi PDF."],
        ["Program", "assessment_programs, assessment_program_simulations", "Periode kegiatan dan simulasi yang dijadwalkan."],
        ["Penugasan", "assessment_participants, assessor_assignments", "Peserta, kategori, pilihan Madya, dan tim asesor."],
        ["Pelaksanaan", "simulation_sessions, simulation_submissions", "Waktu mulai/kedaluwarsa, jawaban, dan file submission."],
        ["Penilaian", "simulation_reviews", "Review per session dan asesor dalam status draft/final."],
        ["Pengawasan", "simulation_session_events, audit_logs", "Event fullscreen dan jejak aktivitas administratif tertentu."],
    ],
)
add_heading(doc, "Persyaratan Server yang Direkomendasikan", level=2)
add_callout(
    doc,
    "Status",
    "Spesifikasi berikut adalah rekomendasi deployment, bukan konfigurasi yang dapat dipastikan dari source code.",
    warning=True,
)
add_table(
    doc,
    ["Parameter", "Rekomendasi Awal"],
    [
        ["CPU", "Minimal 2 vCPU"],
        ["Memori", "Minimal 4 GB RAM"],
        ["Penyimpanan", "Minimal 40 GB SSD dan disesuaikan pertumbuhan PDF/submission"],
        ["Web Server", "Nginx atau Apache dengan HTTPS"],
        ["Runtime", "PHP 8.2+, Composer, Node.js untuk build deployment"],
        ["Database", "PostgreSQL atau MySQL yang didukung Laravel"],
        ["Process", "PHP-FPM dan queue worker yang dikelola supervisor/systemd"],
        ["Backup", "Backup database dan file storage terjadwal serta diuji proses restore-nya"],
    ],
)
add_heading(doc, "Kontrol Keamanan yang Tersedia", level=2)
add_bullets(
    doc,
    [
        "Autentikasi email/password berbasis session dengan regenerasi session setelah login.",
        "Pembatasan percobaan login berdasarkan email dan IP.",
        "CSRF protection dan password hashing bawaan Laravel.",
        "Role middleware serta pemeriksaan akses object pada controller peserta dan asesor.",
        "Sanitasi HTML jawaban rich text sebelum penyimpanan.",
        "Materi PDF disajikan melalui controller dengan authorization dan cache private/no-store.",
        "Audit log untuk aktivitas tertentu dan pencatatan event fullscreen peserta.",
    ],
)

# IV
add_heading(doc, "IV. Fitur Fungsionalitas", level=1)
add_heading(doc, "Status Modul", level=2)
add_table(
    doc,
    ["Modul", "Status", "Keterangan"],
    [
        ["Autentikasi & role", "Aktif", "Login, logout, remember me, redirect dashboard, dan pembatasan route."],
        ["Program Assessment", "Aktif", "CRUD program, periode, status, setup peserta, dan setup asesor."],
        ["Bank Simulasi", "Aktif", "Katalog tetap; admin mengubah materi, deskripsi, dan durasi."],
        ["Peserta Assessment", "Aktif", "Akun manual, pencarian, edit, delete terbatas, dan bulk action."],
        ["Penugasan Asesor", "Aktif", "Satu tim asesor diterapkan untuk seluruh simulasi program."],
        ["Pengerjaan Peserta", "Aktif", "Empat simulasi, timer, PDF preview, rich text, dan upload PDF."],
        ["Monitoring & Penilaian", "Aktif", "Monitoring program, review draft/final, dan hasil peserta."],
        ["Anti-cheat fullscreen", "Parsial", "Mendeteksi keluar fullscreen, penolakan fullscreen, dan perpindahan tab pada alur tertentu."],
        ["Audit log", "Parsial", "Belum seluruh perubahan domain assessment dicatat."],
        ["Integrasi HRIS", "Persiapan", "Schema, konfigurasi, dan indikator UI tersedia; sinkronisasi aktual belum ada."],
        ["SSO INKA", "Rencana", "Belum terdapat protokol, route callback, atau library integrasi."],
        ["Undangan email", "Rencana", "Belum terdapat mailable, token invitation, atau proses distribusi akun."],
        ["Laporan hasil teragregasi", "Rencana", "Hasil masih per review asesor; belum ada konsolidasi dan export final."],
    ],
)
add_heading(doc, "Modul Program Assessment", level=2)
add_bullets(
    doc,
    [
        "Membuat, mengubah, dan menghapus program dengan aturan status tertentu.",
        "Menyimpan nama, deskripsi, periode, status, dan identitas internal program.",
        "Memilih banyak peserta serta kategori assessment dalam satu program.",
        "Memilih minimal satu asesor dan menerapkannya ke seluruh simulasi program.",
        "Memasang katalog simulasi tetap saat setup program disimpan.",
    ],
)
add_heading(doc, "Modul Bank Simulasi", level=2)
add_table(
    doc,
    ["Simulasi", "Materi & Jawaban", "Timer", "Output"],
    [
        ["1. Problem Analysis", "Tepat satu PDF dan satu jawaban rich text", "Ada", "Submission jawaban"],
        ["2. Leaderless Group Discussion", "Review PDF dan jawaban PA secara read-only", "Ada", "Konfirmasi selesai dan observasi asesor"],
        ["3. Critical Incident / In-Tray", "Beberapa PDF; satu jawaban rich text per materi", "Satu timer keseluruhan", "Submission per materi"],
        ["4. Presentasi", "Peserta mengunggah PDF", "Tidak ditemukan timer pengerjaan khusus", "File PDF untuk preview/download asesor"],
    ],
)
add_heading(doc, "Paket Simulasi 3", level=2)
add_table(
    doc,
    ["Paket", "Kategori"],
    [
        ["CI 1", "Kenaikan Golongan I ke II, Promosi Spesialis Pratama, Promosi SPV"],
        ["CI 2", "Kenaikan Golongan II ke III, Promosi Spesialis Muda"],
        ["CI 3", "Kenaikan Golongan III ke IV; opsi pertama untuk Promosi Spesialis Madya"],
        ["In-Tray 1", "Promosi M"],
        ["In-Tray 2", "Promosi SM"],
        ["In-Tray 3", "Opsi kedua untuk Promosi Spesialis Madya"],
    ],
)
add_callout(
    doc,
    "Ketentuan Madya",
    "Asesor memilih CI 3 atau In-Tray 3. Pilihan dikunci ketika salah satu paket telah mulai dikerjakan peserta.",
)
add_heading(doc, "Modul Monitoring dan Penilaian", level=2)
add_bullets(
    doc,
    [
        "Monitoring Admin dan asesor dikelompokkan berdasarkan Program Assessment.",
        "Metrik mencakup peserta, session mulai, submission, review final, dan event fullscreen.",
        "Asesor hanya membuka submission pada program/simulasi yang ditugaskan.",
        "Penilaian dapat disimpan sebagai draft atau difinalisasi.",
        "Peserta hanya melihat review final miliknya sendiri.",
    ],
)

# V
add_heading(doc, "V. Pembagian Role & Tanggung Jawab", level=1)
add_heading(doc, "Detail Hak Akses per Role", level=2)
add_table(
    doc,
    ["Role", "Hak Akses yang Ditemukan", "Batasan"],
    [
        ["Super Admin", "Dashboard, manajemen pengguna, matriks role, audit log, serta fitur administratif Admin HCGA", "Tidak otomatis mengakses workspace operasional Asesor/Peserta"],
        ["Admin HCGA", "Program, Bank Simulasi, peserta, penugasan asesor, dan monitoring", "Tidak memberi penilaian sebagai asesor"],
        ["Asesor", "Simulasi ditugaskan, peserta, monitoring, pilihan Madya, submission, dan review", "Hanya untuk assignment miliknya"],
        ["Peserta Assessment", "Dashboard, simulasi, jadwal, pengerjaan, dan hasil", "Hanya data peserta yang sedang login"],
    ],
)
add_heading(doc, "Matriks Kewenangan", level=2)
add_table(
    doc,
    ["Aktivitas", "Super Admin", "Admin HCGA", "Asesor", "Peserta"],
    [
        ["Kelola akun pengguna", "Ya", "Peserta saja", "Tidak", "Tidak"],
        ["Kelola Program Assessment", "Ya", "Ya", "Tidak", "Tidak"],
        ["Kelola materi simulasi", "Ya", "Ya", "Tidak", "Tidak"],
        ["Atur tim asesor", "Ya", "Ya", "Tidak", "Tidak"],
        ["Pilih paket Madya", "Tidak", "Tidak", "Ya", "Tidak"],
        ["Mengerjakan simulasi", "Tidak", "Tidak", "Tidak", "Ya"],
        ["Memberikan rekomendasi", "Tidak", "Tidak", "Ya", "Tidak"],
        ["Melihat audit log", "Ya", "Tidak", "Tidak", "Tidak"],
    ],
)

# VI
add_heading(doc, "VI. Dokumentasi Antarmuka dan Alur Sistem", level=1)
add_paragraph(
    doc,
    "Bagian ini mendokumentasikan layar yang benar-benar memiliki route aktif. Tangkapan layar visual dapat ditambahkan pada revisi berikutnya setelah environment demonstrasi dan data contoh disepakati. Ketiadaan gambar pada draft ini tidak digunakan untuk mengasumsikan keberadaan fitur lain.",
)
add_heading(doc, "6.1 Halaman Login", level=3)
add_paragraph(doc, "Halaman login menggunakan email, password, opsi Ingat Saya, dan redirect otomatis ke dashboard sesuai role. Tidak terdapat signup aktif, passkey, atau autentikasi dua faktor pada route aplikasi.")
add_heading(doc, "6.2 Dashboard Berdasarkan Role", level=3)
add_paragraph(doc, "Dashboard menampilkan metrik dan chart sesuai role. Komponen ApexCharts dimuat secara dinamis hanya ketika elemen chart tersedia pada halaman.")
add_heading(doc, "6.3 Program Assessment", level=3)
add_paragraph(doc, "Admin melihat daftar program terbaru, melakukan pencarian realtime, memilih beberapa program untuk bulk action, serta membuka fungsi Atur, Edit, atau Hapus sesuai validasi.")
add_heading(doc, "6.4 Setup Program", level=3)
add_paragraph(doc, "Setup program menggabungkan pilihan peserta, kategori assessment, dan tim asesor. Minimal satu asesor wajib dipilih. Seluruh katalog simulasi tetap dipasang oleh backend.")
add_heading(doc, "6.5 Bank Simulasi", level=3)
add_paragraph(doc, "Admin mengubah materi dan durasi pada simulasi yang sudah ditetapkan. Tidak tersedia tombol membuat atau menghapus jenis simulasi.")
add_heading(doc, "6.6 Simulasi Saya", level=3)
add_paragraph(doc, "Peserta melihat simulasi yang masih relevan. Simulasi yang telah dikumpulkan atau melewati batas waktu disembunyikan dari daftar aktif.")
add_heading(doc, "6.7 Halaman Pengerjaan", level=3)
add_paragraph(doc, "Halaman pengerjaan menampilkan timer server-side, preview PDF, rich text editor, serta kontrol simpan/kumpulkan. Pada Simulasi 3, jawaban antarmateri disimpan tanpa meninggalkan fullscreen.")
add_heading(doc, "6.8 Monitoring", level=3)
add_paragraph(doc, "Monitoring Admin dan asesor menampilkan progres pada tingkat program dan detail peserta. Event fullscreen berfungsi sebagai konteks pengawasan, bukan keputusan otomatis.")
add_heading(doc, "6.9 Penilaian & Rekomendasi", level=3)
add_paragraph(doc, "Asesor meninjau materi dan submission, menulis catatan, memilih rekomendasi, lalu menyimpan draft atau final. Penilaian final dikunci dari alur normal.")
add_heading(doc, "6.10 Hasil Peserta", level=3)
add_paragraph(doc, "Peserta melihat rekomendasi dan catatan yang sudah difinalisasi. Source code belum menyediakan satu keputusan agregat lintas asesor maupun export laporan akhir.")

add_heading(doc, "Alur End-to-End", level=2)
add_steps(
    doc,
    [
        "Super Admin memastikan akun Admin HCGA dan asesor aktif.",
        "Admin HCGA menyiapkan akun peserta dan memperbarui Bank Simulasi.",
        "Admin membuat Program Assessment beserta periode dan status.",
        "Admin memilih peserta, kategori assessment, dan tim asesor pada setup program.",
        "Peserta membuka jadwal, memulai session, dan menyelesaikan simulasi yang relevan.",
        "Admin dan asesor memonitor progres serta event pelaksanaan.",
        "Asesor memeriksa submission, menyimpan draft, dan memfinalisasi rekomendasi.",
        "Peserta melihat hasil yang telah difinalisasi.",
    ],
)
add_heading(doc, "Alur Status Session", level=2)
add_table(
    doc,
    ["Tahap", "Status/Kondisi", "Perilaku Sistem"],
    [
        ["Belum dimulai", "Tidak ada session aktif", "Peserta dapat memulai jika program dan jadwal tersedia."],
        ["Sedang dikerjakan", "in_progress", "Timer memakai expires_at server dan session yang sama digunakan kembali."],
        ["Waktu habis", "expired", "Session ditandai kedaluwarsa ketika diperiksa setelah batas waktu."],
        ["Dikumpulkan", "submitted", "Submission tersimpan dan simulasi tidak dapat diedit melalui alur normal."],
        ["Dinilai", "Review draft/submitted", "Draft hanya untuk asesor; final dapat ditampilkan kepada peserta."],
    ],
)

add_heading(doc, "Roadmap Integrasi", level=2)
add_table(
    doc,
    ["Tahap", "Ruang Lingkup", "Prasyarat dari PT INKA"],
    [
        ["1. Finalisasi Assessment", "Perbaikan konflik kode program, audit, validasi delete, test, dan laporan final", "Persetujuan proses bisnis dan format output"],
        ["2. Integrasi HRIS", "Sinkronisasi identitas dan status pegawai", "Dokumentasi API, sandbox, credential, identifier unik, dan aturan data"],
        ["3. SSO INKA", "Login melalui identity provider perusahaan", "Protokol OIDC/SAML, client ID, secret/certificate, callback, claim mapping"],
        ["4. Undangan Email", "Undangan aman dan aktivasi akses peserta", "SMTP/API email, template, sender domain, dan kebijakan masa berlaku"],
        ["5. Laporan Akhir", "Konsolidasi review, approval, dan export", "Format laporan resmi dan aturan keputusan"],
    ],
)

add_heading(doc, "Catatan Risiko dan Gap", level=2)
add_table(
    doc,
    ["Temuan", "Dampak", "Rekomendasi"],
    [
        ["Kode internal program masih tampil pada satu halaman", "Pengguna melihat kode yang tidak digunakan proses bisnis", "Sembunyikan dari seluruh UI; pertahankan hanya sebagai identifier internal bila diperlukan"],
        ["Audit log belum mencakup seluruh perubahan", "Jejak perubahan belum lengkap", "Tambahkan pencatatan program, katalog, setup, submission penting, dan finalisasi review"],
        ["Delete program draft belum memeriksa seluruh riwayat", "Potensi cascade delete data turunan", "Blokir delete jika sudah memiliki session/submission/review"],
        ["Super Admin bukan full access lintas role", "Makna full access dapat berbeda dengan ekspektasi", "Tetapkan apakah membutuhkan impersonation atau cukup akses administratif"],
        ["Anti-cheat tidak seragam pada Presentasi", "Cakupan monitoring berbeda antar-simulasi", "Putuskan apakah upload presentasi memang memerlukan fullscreen"],
        ["Hasil belum teragregasi", "Belum ada satu keputusan final tingkat program", "Definisikan workflow konsolidasi, approval, dan export"],
        ["Seeder memakai password contoh", "Risiko jika dijalankan di production", "Pisahkan seeder demo dan production serta wajibkan rotasi password"],
    ],
)

# VII
add_heading(doc, "VII. Kesimpulan", level=1)
add_paragraph(
    doc,
    "INKA Assessment Management System telah memiliki fondasi operasional assessment yang terhubung dari persiapan program hingga penyampaian hasil final. Implementasi saat ini mencakup empat role, katalog empat simulasi, pengelolaan materi PDF, penugasan tim asesor, session bertimer, rich text response, monitoring, penilaian, serta kontrol akses berbasis role dan kepemilikan data.",
)
add_paragraph(
    doc,
    "Aplikasi layak dilanjutkan menuju tahap finalisasi dan acceptance testing modul Assessment. Sebelum ditetapkan sebagai versi produksi final, prioritas berikutnya adalah menyelesaikan gap validasi dan audit, menyepakati bentuk laporan hasil teragregasi, membersihkan artefak template/legacy, serta memastikan seluruh automated test kembali lulus.",
)
add_paragraph(
    doc,
    "Integrasi HRIS, SSO INKA, dan undangan email ditempatkan sebagai roadmap terpisah karena source code baru menyediakan sebagian fondasi data dan konfigurasi. Implementasinya memerlukan dokumentasi teknis, credential, serta keputusan proses bisnis resmi dari PT INKA (Persero).",
)

doc.core_properties.title = "Proposal & Blueprint INKA Assessment Management System"
doc.core_properties.subject = "Proposal dan blueprint teknis modul Assessment"
doc.core_properties.author = "PT INKA (Persero) - Divisi Human Capital"
doc.core_properties.keywords = "INKA, assessment, proposal, blueprint, Human Capital"
doc.core_properties.comments = "Disusun berdasarkan source code assessment-inka-final per 7 September 2026."

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
